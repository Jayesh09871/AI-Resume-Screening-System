import os
import io
import re
from typing import Dict, List, Tuple, Any
import fitz  # PyMuPDF
from docx import Document
from backend.app.utils.text_cleaner import normalize_whitespace
from backend.app.utils.logger import logger

SECTION_PATTERNS = {
    "Summary": [r"^summary\b", r"^professional summary\b", r"^profile\b", r"^about me\b", r"^objective\b"],
    "Skills": [r"^skills\b", r"^technical skills\b", r"^core competencies\b", r"^technologies\b", r"^key skills\b"],
    "Experience": [r"^experience\b", r"^work experience\b", r"^professional experience\b", r"^employment\b", r"^work history\b"],
    "Education": [r"^education\b", r"^academic background\b", r"^academics\b", r"^qualifications\b"],
    "Projects": [r"^projects\b", r"^personal projects\b", r"^key projects\b", r"^technical projects\b"],
    "Certifications": [r"^certifications\b", r"^certificates\b", r"^courses\b", r"^licenses\b"],
    "Achievements": [r"^achievements\b", r"^awards\b", r"^honors\b", r"^accomplishments\b"],
    "Links": [r"^links\b", r"^profiles\b", r"^socials\b", r"^portfolio\b"],
}


class ResumeParser:
    """
    Parses resume files (PDF or DOCX) into clean text, page counts,
    and identifies major resume sections with safety validation.
    """

    @staticmethod
    def validate_file(filename: str, file_size: int, content_type: str, max_size_mb: int = 10) -> List[str]:
        warnings = []
        ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
        if ext not in ["pdf", "docx"]:
            raise ValueError(f"Unsupported file format '.{ext}'. Only PDF and DOCX files are allowed.")
        
        max_bytes = max_size_mb * 1024 * 1024
        if file_size <= 0:
            raise ValueError("The uploaded file is empty.")
        if file_size > max_bytes:
            raise ValueError(f"File size ({round(file_size / (1024 * 1024), 2)}MB) exceeds maximum allowed limit of {max_size_mb}MB.")
        
        return warnings

    @classmethod
    def parse_pdf(cls, file_bytes: bytes) -> Tuple[str, int, List[str]]:
        """Extract text from PDF using PyMuPDF page-by-page."""
        warnings = []
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
        except Exception as e:
            logger.error(f"Failed to open PDF file: {e}")
            raise ValueError("The uploaded PDF file is corrupted or cannot be read.")

        page_count = len(doc)
        if page_count == 0:
            raise ValueError("The uploaded PDF contains no pages.")

        extracted_pages = []
        for i in range(page_count):
            try:
                page = doc.load_page(i)
                text = page.get_text("text")
                if text:
                    extracted_pages.append(text)
            except Exception as e:
                warnings.append(f"Warning: Could not read text from page {i + 1}: {str(e)}")

        full_text = "\n\n".join(extracted_pages)
        if not full_text.strip():
            warnings.append("Notice: Document contains no extractable text. It might be scanned or image-only.")

        return normalize_whitespace(full_text), page_count, warnings

    @classmethod
    def parse_docx(cls, file_bytes: bytes) -> Tuple[str, int, List[str]]:
        """Extract text from DOCX using python-docx."""
        warnings = []
        try:
            doc = Document(io.BytesIO(file_bytes))
        except Exception as e:
            logger.error(f"Failed to open DOCX file: {e}")
            raise ValueError("The uploaded DOCX file is corrupted or cannot be read.")

        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text if present
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    table_texts.append(row_text)

        full_content = "\n".join(paragraphs + table_texts)
        if not full_content.strip():
            warnings.append("Notice: DOCX file does not contain readable paragraph or table text.")

        # In docx, approximate page count (every ~350 words is ~1 page)
        words = len(full_content.split())
        approx_pages = max(1, (words // 350) + (1 if words % 350 > 0 else 0))

        return normalize_whitespace(full_content), approx_pages, warnings

    @classmethod
    def detect_sections(cls, text: str) -> List[str]:
        """Detect which key resume sections are present in the text."""
        detected = []
        lines = text.split("\n")
        
        for section_name, patterns in SECTION_PATTERNS.items():
            for line in lines:
                clean_line = line.strip().lower()
                # Check line against patterns
                matched = any(re.search(pat, clean_line) for pat in patterns)
                if matched:
                    if section_name not in detected:
                        detected.append(section_name)
                    break

        return detected

    @classmethod
    def parse_document(cls, filename: str, file_bytes: bytes, content_type: str = "") -> Dict[str, Any]:
        """Main entry point to validate and parse resume document."""
        cls.validate_file(filename, len(file_bytes), content_type)
        ext = filename.rsplit(".", 1)[-1].lower()

        if ext == "pdf":
            raw_text, page_count, warnings = cls.parse_pdf(file_bytes)
        elif ext == "docx":
            raw_text, page_count, warnings = cls.parse_docx(file_bytes)
        else:
            raise ValueError(f"Unsupported format: {ext}")

        detected_sections = cls.detect_sections(raw_text)

        return {
            "raw_text": raw_text,
            "page_count": page_count,
            "detected_sections": detected_sections,
            "warnings": warnings,
        }
