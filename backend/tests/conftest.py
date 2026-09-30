import io
import pytest
from fastapi.testclient import TestClient
import fitz
from docx import Document
from backend.app.main import app
from backend.app.database.database import Base, engine, SessionLocal
from backend.app.models.schemas import ResumeSchema, ExperienceItem, EducationItem, ProjectItem, JobDescriptionSchema

SAMPLE_RESUME_TEXT = """
ALEXANDER MORGAN
alex.morgan.fake@example.com | (555) 234-5678 | San Francisco, CA | https://github.com/alexmorgan-fake | https://linkedin.com/in/alexmorgan-fake

PROFESSIONAL SUMMARY
Results-oriented Software Engineer with 4+ years of experience designing and deploying high-performance backend systems, distributed microservices, and RESTful APIs using Python, FastAPI, and PostgreSQL. Experienced with Docker containerization and Kubernetes orchestration.

SKILLS
Programming Languages: Python, JavaScript, TypeScript, SQL, Go
Frameworks & Libraries: FastAPI, Django, React, Node.js, Express
Databases: PostgreSQL, Redis, MongoDB
DevOps & Cloud: Docker, Kubernetes, AWS, GitHub Actions, CI/CD

PROFESSIONAL EXPERIENCE
Senior Backend Engineer | CloudScale Tech | San Francisco, CA
June 2022 - Present
- Architected and scaled high-throughput microservices using FastAPI and PostgreSQL, serving 2M+ requests daily.
- Implemented asynchronous background task processing using Redis and Celery, reducing API response times by 35%.
- Built automated CI/CD deployment pipelines utilizing Docker and GitHub Actions for zero-downtime releases.

Software Engineer | NextGen Innovations | Austin, TX
July 2020 - May 2022
- Developed RESTful APIs using Python and Django for an enterprise SaaS platform.
- Managed PostgreSQL database schema migrations and query optimizations.
- Collaborated in an agile scrum team to deliver product features and unit test suites.

EDUCATION
Bachelor of Science in Computer Science | University of California, Berkeley | 2020
GPA: 3.8 / 4.0

PROJECTS
Real-Time Distributed Queue
- Built a fault-tolerant message queue in Python and Go using Redis and WebSockets.
- Created interactive React monitoring dashboard for message consumer health metrics.

CERTIFICATIONS
- AWS Certified Solutions Architect - Associate
"""

SAMPLE_JD_TEXT = """
Senior Python / Backend Developer
TechForward Inc. - Remote / San Francisco, CA

About the Role:
We are seeking an experienced Backend Software Engineer to build resilient distributed systems and API architectures.

Key Responsibilities:
- Design, build, and maintain high-volume RESTful APIs and microservices in Python.
- Optimize database schemas and queries in PostgreSQL.
- Implement containerized deployments using Docker and Kubernetes.
- Collaborate with frontend engineers using React to integrate scalable endpoints.

Required Qualifications & Skills:
- 3+ years of professional backend software engineering experience.
- Strong proficiency in Python and modern frameworks like FastAPI or Django.
- Solid experience with relational databases, specifically PostgreSQL.
- Practical experience with Docker and REST APIs.

Preferred Qualifications:
- Experience with Redis caching and asynchronous queues.
- Familiarity with Kubernetes, AWS cloud infrastructure, and CI/CD pipelines.
- Bachelor's degree in Computer Science or related STEM field.
"""


@pytest.fixture(scope="session")
def client():
    # Setup test DB tables
    Base.metadata.create_all(bind=engine)
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture
def db_session():
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


@pytest.fixture
def sample_resume_text():
    return SAMPLE_RESUME_TEXT


@pytest.fixture
def sample_jd_text():
    return SAMPLE_JD_TEXT


@pytest.fixture
def sample_resume_model():
    return ResumeSchema(
        name="Alexander Morgan",
        email="alex.morgan.fake@example.com",
        phone="(555) 234-5678",
        location="San Francisco, CA",
        summary="Results-oriented Software Engineer with 4+ years experience designing backend systems using Python, FastAPI, and PostgreSQL.",
        skills=["Python", "FastAPI", "PostgreSQL", "Docker", "Redis", "React", "AWS", "Git"],
        experience=[
            ExperienceItem(
                title="Senior Backend Engineer",
                company="CloudScale Tech",
                location="San Francisco, CA",
                start_date="2022",
                end_date="Present",
                highlights=[
                    "Architected and scaled microservices using FastAPI and PostgreSQL.",
                    "Implemented background queues with Redis, improving latency by 35%."
                ]
            )
        ],
        education=[
            EducationItem(
                degree="B.S. in Computer Science",
                institution="UC Berkeley",
                graduation_year="2020",
                gpa="3.8"
            )
        ],
        projects=[
            ProjectItem(
                name="Distributed Queue",
                description="Message streaming system",
                technologies=["Python", "Redis", "Docker"],
                link="https://github.com/example/queue",
                highlights=["Engineered reliable queuing mechanisms."]
            )
        ],
        certifications=["AWS Certified Solutions Architect"],
        achievements=["Dean's Honors"],
        links=["https://github.com/alexmorgan-fake", "https://linkedin.com/in/alexmorgan-fake"]
    )


@pytest.fixture
def sample_pdf_bytes():
    """Create a valid in-memory PDF file for testing."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), SAMPLE_RESUME_TEXT)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_docx_bytes():
    """Create a valid in-memory DOCX file for testing."""
    doc = Document()
    doc.add_heading("Alexander Morgan", 0)
    for line in SAMPLE_RESUME_TEXT.split("\n"):
        if line.strip():
            doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
